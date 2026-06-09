"""Unit tests for SCOUT pipeline.

Run with:  python -m unittest discover -s tests
"""
from __future__ import annotations

import sys
import unittest
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parent.parent))

from src.criteria_builder import (
    HardFilter, build_boolean_query_dict, generate_boolean_query,
    load_guidelines,
)
from src.jd_loader import (
    JobDescription, _read_excel, _parse_range_string, _truthy,
    load_jds,
)


def _settings():
    from src.config_loader import load_settings
    return load_settings()


def _jd(skills=("Python", "FastAPI"), exp_min=5, exp_max=8, locs=("Chennai",),
        ctc=25.0, fresh=45):
    return JobDescription(
        jd_id="JD001", role="Python Developer", skills=list(skills),
        experience_min=exp_min, experience_max=exp_max, locations=list(locs),
        notes="", ctc_ceiling_lacs=ctc, freshness_days=fresh,
    )


class TestStaticBooleanQuery(unittest.TestCase):
    def test_joins_with_and(self):
        jd = _jd()
        self.assertEqual(build_boolean_query_dict(jd.to_dict()), '"Python" AND "FastAPI"')

    def test_empty_skills(self):
        jd = _jd(skills=())
        self.assertEqual(build_boolean_query_dict(jd.to_dict()), "")


class TestClaudeBooleanQuery(unittest.TestCase):
    def test_generates_with_claude_or_falls_back(self):
        s = _settings()
        jd = _jd()
        result = generate_boolean_query(s, jd.to_dict(), hint=None, attempt=1)
        self.assertIsInstance(result, str)
        self.assertGreater(len(result), 0)
        print(f"\n[BooleanQuery] '{result}'")

    def test_max_attempts_raises(self):
        s = _settings()
        jd = _jd()
        with self.assertRaises(ValueError):
            generate_boolean_query(s, jd.to_dict(), attempt=99)


class TestHardFilter(unittest.TestCase):
    def _f(self, **kw):
        defaults = dict(
            mandatory_skills=["Python", "FastAPI"],
            experience_min=5, experience_max=8,
            locations=["Chennai", "Bangalore"],
            max_notice_period_days=60,
            ctc_ceiling_lacs=25.0,
            freshness_days=45,
            min_ai_score=70,
        )
        defaults.update(kw)
        return HardFilter(**defaults)

    def test_accepts_good_candidate(self):
        rej, _ = self._f().rejects({
            "skills": ["Python", "FastAPI", "AWS"],
            "experience_years": 6, "location": "Chennai",
            "notice_period_days": 30, "current_ctc_lacs": 22.0,
            "last_active_days": 10,
        })
        self.assertFalse(rej)

    def test_rejects_missing_mandatory(self):
        rej, _ = self._f().rejects({
            "skills": ["Python", "Django"],
            "experience_years": 6, "location": "Chennai",
        })
        self.assertTrue(rej)

    def test_rejects_ctc_above_ceiling(self):
        rej, _ = self._f().rejects({
            "skills": ["Python", "FastAPI"],
            "experience_years": 6, "location": "Chennai",
            "current_ctc_lacs": 30.0,
        })
        self.assertTrue(rej)

    def test_rejects_stale_profile(self):
        rej, _ = self._f().rejects({
            "skills": ["Python", "FastAPI"],
            "experience_years": 6, "location": "Chennai",
            "last_active_days": 100,
        })
        self.assertTrue(rej)

    def test_rejects_long_notice_period(self):
        rej, _ = self._f().rejects({
            "skills": ["Python", "FastAPI"],
            "experience_years": 6, "location": "Chennai",
            "notice_period_days": 90,
        })
        self.assertTrue(rej)

    def test_rejects_wrong_location(self):
        rej, _ = self._f().rejects({
            "skills": ["Python", "FastAPI"],
            "experience_years": 6, "location": "Delhi",
        })
        self.assertTrue(rej)

    def test_rejects_below_min_experience(self):
        rej, _ = self._f().rejects({
            "skills": ["Python", "FastAPI"],
            "experience_years": 3, "location": "Chennai",
        })
        self.assertTrue(rej)

    def test_rejects_above_max_experience(self):
        rej, _ = self._f().rejects({
            "skills": ["Python", "FastAPI"],
            "experience_years": 12, "location": "Chennai",
        })
        self.assertTrue(rej)

    def test_skips_ctc_when_unset(self):
        f = self._f(ctc_ceiling_lacs=None)
        rej, _ = f.rejects({
            "skills": ["Python", "FastAPI"],
            "experience_years": 6, "location": "Chennai",
            "current_ctc_lacs": 100.0,
        })
        self.assertFalse(rej)


class TestDedup(unittest.TestCase):
    """Dedup key = (name.lower(), company.lower())."""

    def test_dedup_set_behavior(self):
        seen: set[tuple[str, str]] = set()
        for raw in [("Alice Kumar", "Acme Corp"),
                    ("alice kumar", "acme corp"),
                    ("ALICE KUMAR", "ACME CORP")]:
            seen.add((raw[0].lower(), raw[1].lower()))
        self.assertEqual(len(seen), 1)

    def test_different_company_not_deduped(self):
        seen: set[tuple[str, str]] = set()
        seen.add(("alice kumar", "acme corp"))
        seen.add(("alice kumar", "beta inc"))
        self.assertEqual(len(seen), 2)

    def test_empty_name_deduped(self):
        seen: set[tuple[str, str]] = set()
        for _ in range(3):
            seen.add(("", ""))
        self.assertEqual(len(seen), 1)


class TestResumeParser(unittest.TestCase):
    def test_extracts_skills(self):
        import fitz
        from src.resume_parser import parse
        doc = fitz.open()
        page = doc.new_page()
        page.insert_text((50, 50),
            "John Doe\n\nSkills\nPython, FastAPI, AWS, Docker\n\nExperience\nX 2020 - Present\n")
        p = Path("/tmp/_test_resume.pdf")
        doc.save(str(p)); doc.close()
        result = parse(p)
        self.assertIn("Python", result["skills"])
        self.assertIn("FastAPI", result["skills"])
        self.assertGreater(len(result["experience"]), 0)


class TestJDLoaderRobustness(unittest.TestCase):
    def _settings(self):
        return _settings()

    def _write(self, name, header_row, headers, rows):
        from openpyxl import Workbook
        wb = Workbook()
        ws = wb.active
        for r in range(1, header_row):
            ws.cell(row=r, column=1, value=f"Title block line {r}")
        for c, h in enumerate(headers, 1):
            ws.cell(row=header_row + 1, column=c, value=h)
        for r_idx, row in enumerate(rows, start=header_row + 2):
            for c_idx, val in enumerate(row, start=1):
                ws.cell(row=r_idx, column=c_idx, value=val)
        path = Path(f"/tmp/{name}")
        wb.save(path)
        return path

    def test_friendly_headers_with_new_scout_columns(self):
        path = self._write("scout_friendly.xlsx", 1, [
            "JD ID", "Role", "Skills Required", "Min Years", "Max Years",
            "Location", "CTC Ceiling", "Freshness",
            "Status", "Hiring Manager"
        ], [
            ["JD001", "Python Developer", "Python, FastAPI", 5, 8,
             "Chennai", 25, 45, "Active", "Priya"],
        ])
        result = load_jds(self._settings(), excel_path=path)
        self.assertEqual(len(result), 1)
        self.assertEqual(result[0].ctc_ceiling_lacs, 25.0)
        self.assertEqual(result[0].freshness_days, 45)

    def test_excludes_inactive_rows(self):
        path = self._write("active.xlsx", 1, [
            "JD_ID", "Role", "Skills", "Experience_Min_Years",
            "Experience_Max_Years", "Location", "Active"
        ], [
            ["JD020", "A", "X", 1, 2, "Chennai", "TRUE"],
            ["JD021", "B", "Y", 1, 2, "Chennai", "Closed"],
        ])
        result = load_jds(self._settings(), excel_path=path)
        self.assertEqual([j.jd_id for j in result], ["JD020"])


class TestParseRangeString(unittest.TestCase):
    """Test parsing of combined experience range strings."""

    def test_standard_range(self):
        self.assertEqual(_parse_range_string("5-8 yrs"), (5, 8))

    def test_range_with_plus(self):
        self.assertEqual(_parse_range_string("5+ to 8+ Years"), (5, 8))

    def test_range_with_spaces(self):
        self.assertEqual(_parse_range_string("5 - 8 years"), (5, 8))

    def test_single_number(self):
        self.assertEqual(_parse_range_string("5 years"), (5, 5))
        self.assertEqual(_parse_range_string("5+ yrs"), (5, 5))

    def test_empty_string(self):
        self.assertEqual(_parse_range_string(""), (0, 0))
        self.assertEqual(_parse_range_string(None), (0, 0))

    def test_non_numeric(self):
        self.assertEqual(_parse_range_string("Fresher"), (0, 0))


class TestTruthyWithStatusValues(unittest.TestCase):
    """The _truthy function must accept Naukri Search Automation statuses."""

    def test_pending_is_truthy(self):
        self.assertTrue(_truthy("Pending"))

    def test_in_progress_is_truthy(self):
        self.assertTrue(_truthy("In Progress"))

    def test_standard_truthy_values(self):
        self.assertTrue(_truthy(True))
        self.assertTrue(_truthy("TRUE"))
        self.assertTrue(_truthy("Active"))
        self.assertTrue(_truthy("Open"))

    def test_falsy_values(self):
        self.assertFalse(_truthy(False))
        self.assertFalse(_truthy("Closed"))
        self.assertFalse(_truthy("Completed"))
        self.assertFalse(_truthy(""))
        self.assertFalse(_truthy(None))


class TestNewExcelFormat(unittest.TestCase):
    """Test that the Naukri Search Automation Excel format loads correctly."""

    def _settings(self):
        return _settings()

    def _write(self, name, headers, rows):
        from openpyxl import Workbook
        wb = Workbook()
        ws = wb.active
        for c, h in enumerate(headers, 1):
            ws.cell(row=1, column=c, value=h)
        for r_idx, row in enumerate(rows, start=2):
            for c_idx, val in enumerate(row, start=1):
                ws.cell(row=r_idx, column=c_idx, value=val)
        path = Path(f"/tmp/{name}")
        wb.save(path)
        return path

    def test_loads_new_format_with_jd_description(self):
        path = self._write("new_format.xlsx", [
            "Job ID", "Job Title", "Job Description", "Location",
            "Experience Required", "Key Skills", "Priority", "Status",
        ], [
            ["JD-001", "Data Scientist",
             "Role: Data Scientist\nSkills: Python, ML\nExperience: 5+ yrs",
             "Remote / India", "5-8 yrs",
             "Python, Machine Learning, SQL", "High", "Pending"],
        ])
        result = load_jds(self._settings(), excel_path=path)
        self.assertEqual(len(result), 1)
        jd = result[0]
        self.assertEqual(jd.jd_id, "JD-001")
        self.assertEqual(jd.role, "Data Scientist")
        self.assertIn("Python", jd.skills)
        self.assertIn("Machine Learning", jd.skills)
        self.assertEqual(jd.experience_min, 5)
        self.assertEqual(jd.experience_max, 8)
        self.assertIn("Remote", jd.locations[0])
        # jd_description should be populated
        self.assertIn("Role: Data Scientist", jd.jd_description)
        self.assertIn("Skills: Python", jd.jd_description)

    def test_loads_new_format_inactive_status(self):
        path = self._write("inactive.xlsx", [
            "Job ID", "Job Title", "Job Description", "Location",
            "Experience Required", "Key Skills", "Status",
        ], [
            ["JD-010", "Role A", "desc", "India", "3-5 yrs", "SkillX", "Pending"],
            ["JD-011", "Role B", "desc", "India", "3-5 yrs", "SkillY", "Completed"],
            ["JD-012", "Role C", "desc", "India", "3-5 yrs", "SkillZ", "Closed"],
        ])
        result = load_jds(self._settings(), excel_path=path)
        self.assertEqual([j.jd_id for j in result], ["JD-010"])

    def test_missing_ctc_and_freshness_uses_none(self):
        """New format doesn't have CTC/Freshness columns; should default to None."""
        path = self._write("no_ctc.xlsx", [
            "Job ID", "Job Title", "Job Description", "Location",
            "Experience Required", "Key Skills", "Status",
        ], [
            ["JD-020", "Engineer", "desc", "India", "5-8 yrs", "SkillA", "Pending"],
        ])
        result = load_jds(self._settings(), excel_path=path)
        jd = result[0]
        self.assertIsNone(jd.ctc_ceiling_lacs)
        self.assertIsNone(jd.freshness_days)

    def test_jd_description_in_to_dict(self):
        jd = JobDescription(
            jd_id="JD-TEST", role="Tester", skills=["X"],
            experience_min=1, experience_max=5, locations=["Any"],
            notes="", jd_description="Full JD text here",
        )
        d = jd.to_dict()
        self.assertEqual(d["jd_description"], "Full JD text here")


class TestSCOUTFiltersInHardFilter(unittest.TestCase):
    """The HardFilter must now accept CTC and freshness inputs."""

    def test_ctc_field_default_none(self):
        f = HardFilter()
        self.assertIsNone(f.ctc_ceiling_lacs)

    def test_freshness_field_default(self):
        f = HardFilter()
        self.assertIsInstance(f.freshness_days, int)


if __name__ == "__main__":
    unittest.main()
